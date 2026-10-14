"""独立校验 WizDeck 导出的 .pptx：结构、尺寸、形状类型、文字内容与备注。

用 python-pptx 打开（它对 OOXML 结构敏感，能打开就说明包基本正确），
再逐页打印形状清单，用于交叉验证「文字是矢量、渐变是贴图」这一分层策略。
"""
import sys
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

PPTX = Path(sys.argv[1] if len(sys.argv) > 1 else "tmp/verify/wizdeck-export.pptx")

if not PPTX.exists():
    print(f"找不到文件: {PPTX}")
    sys.exit(1)

print("=" * 78)
print("WizDeck PPTX 校验:", PPTX)
print("=" * 78)

# ── 0. 先确认是个合法 zip 包 ────────────────────────────────────────────
with zipfile.ZipFile(PPTX) as z:
    bad = z.testzip()
    names = z.namelist()
print(f"[zip] 条目 {len(names)} 个, CRC 校验: {'全部通过' if bad is None else '损坏 -> ' + str(bad)}")
required = [
    "[Content_Types].xml",
    "_rels/.rels",
    "ppt/presentation.xml",
    "ppt/slideMasters/slideMaster1.xml",
    "ppt/slideLayouts/slideLayout1.xml",
    "ppt/theme/theme1.xml",
]
for r in required:
    print(f"  {'OK ' if r in names else '缺失'} {r}")

# ── 1. 用 python-pptx 打开 ─────────────────────────────────────────────
prs = Presentation(str(PPTX))
print()
print(f"[尺寸] {prs.slide_width} x {prs.slide_height} EMU  =  "
      f"{Emu(prs.slide_width).inches:.3f}in x {Emu(prs.slide_height).inches:.3f}in")
print(f"       期望 20.000in x 11.250in (1920x1080 px @96dpi 等比)")
print(f"[页数] {len(prs.slides)}")

notes_master = prs.notes_master
print(f"[备注母版] {'存在' if notes_master is not None else '不存在'}")

total_shapes = 0
total_text = 0
total_pics = 0

for i, slide in enumerate(prs.slides, 1):
    print()
    print(f"── 第 {i} 页 " + "─" * 60)
    kinds = {}
    texts = []
    for sh in slide.shapes:
        total_shapes += 1
        t = str(sh.shape_type)
        kinds[t] = kinds.get(t, 0) + 1
        if sh.shape_type == 13:  # PICTURE
            total_pics += 1
        if sh.has_text_frame:
            txt = sh.text_frame.text.strip()
            if txt:
                total_text += 1
                texts.append(txt.replace("\n", " ⏎ "))
    print("   形状构成:", ", ".join(f"{k}×{v}" for k, v in sorted(kinds.items())))
    if texts:
        print(f"   文本框 {len(texts)} 个, 抽样前 4 条:")
        for t in texts[:4]:
            print("     ·", (t[:72] + "…") if len(t) > 72 else t)
    note = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
    print(f"   备注: {note[:60] if note else '（空）'}")

print()
print("=" * 78)
print(f"合计: {len(prs.slides)} 页 / {total_shapes} 个形状 / {total_text} 个带文字形状 / {total_pics} 张图片")
print("=" * 78)

# ── 2. 交叉检查：文字是否真的是可编辑文本 ───────────────────────────────
expect_tokens = ["分数", "学习目标", "单位", "性质"]
all_text = "\n".join(
    sh.text_frame.text
    for s in prs.slides
    for sh in s.shapes
    if sh.has_text_frame
)
print()
print("[文字矢量性检查] 导出的文本框里应能检索到关键内容：")
for tok in expect_tokens:
    print(f"  {'命中' if tok in all_text else '未命中'}  {tok}")
print(f"  文本总长度: {len(all_text)} 字")
