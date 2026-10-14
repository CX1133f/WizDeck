"""PPTX 越界审计：找出任何超出幻灯片边界（或异常小）的形状。

我用 20x11.25in = 1920x1080px 的画布，px→EMU = 9525。
文本框允许换行时宽度会加 10% 缓冲，因此必须确认没有把内容推出可视区。
"""
import sys
import zipfile

from pptx import Presentation

sys.stdout.reconfigure(encoding="utf-8")

PX = 9525
TOL = 2  # 允许 2px 的取整误差

# 这些形状越界是设计使然：装饰性光晕/色块故意超出画布，由画布的 overflow:hidden 裁掉。
# 它们自身不含文字（文字另有独立文本框，且会被夹在画布内），所以不影响可读性。
DECORATIVE_NAMES = {"容器", "区块", "页面底色", "矩形", "椭圆"}


def is_decorative(name, has_text):
    return (not has_text) and (name in DECORATIVE_NAMES)


def audit(path, slide_w_px=1920, slide_h_px=1080):
    prs = Presentation(path)
    W, H = slide_w_px, slide_h_px
    problems = []
    clipped_ok = 0
    total = 0
    for i, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            total += 1
            try:
                l = sh.left / PX
                t = sh.top / PX
                w = sh.width / PX
                h = sh.height / PX
            except TypeError:
                continue
            name = sh.name
            has_text = sh.has_text_frame and bool(sh.text_frame.text.strip())
            issues = []
            if l < -TOL or t < -TOL:
                issues.append(("起点为负", round(l, 1), round(t, 1)))
            if l + w > W + TOL:
                issues.append(("右边越界", round(l + w, 1), ">", W))
            if t + h > H + TOL:
                issues.append(("底边越界", round(t + h, 1), ">", H))
            if w < 1 or h < 1:
                issues.append(("尺寸异常", round(w, 1), round(h, 1)))
            if not issues:
                continue
            if is_decorative(name, has_text):
                clipped_ok += 1
                continue
            for it in issues:
                problems.append((i, name, it[0]) + tuple(it[1:]))
    return total, problems, len(prs.slides), clipped_ok


if __name__ == "__main__":
    paths = sys.argv[1:]
    if not paths:
        paths = ["samples/output/sample-fractions.pptx"]
    grand = 0
    for p in paths:
        if not zipfile.ZipFile(p).testzip() is None:
            print(p, "CRC 损坏")
        total, problems, n, clipped = audit(p)
        grand += len(problems)
        status = "OK" if not problems else "%d 处越界" % len(problems)
        if clipped:
            status += "（%d 个装饰层越界但被画布裁切，可接受）" % clipped
        print("%-46s %2d 页 %4d 形状  %s" % (p.split("/")[-1].split("\\")[-1], n, total, status))
        for pr in problems[:8]:
            print("    第 %d 页 %-16s %s  %s" % (pr[0], pr[1], pr[2], " ".join(str(x) for x in pr[3:])))
        if len(problems) > 8:
            print("    ...以及另外 %d 处" % (len(problems) - 8))
    print()
    print("合计问题:", grand)
    sys.exit(1 if grand else 0)
