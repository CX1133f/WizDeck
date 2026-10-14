"""校验参考产物：结构、尺寸、矢量文字数量。"""
import sys
import zipfile

from pptx import Presentation

sys.stdout.reconfigure(encoding="utf-8")

P = "samples/output/sample-fractions.pptx"
print("CRC:", "OK" if zipfile.ZipFile(P).testzip() is None else "BAD")
prs = Presentation(P)
print("slides:", len(prs.slides),
      "| size:", round(prs.slide_width / 914400, 3), "x", round(prs.slide_height / 914400, 3), "in")
text_shapes = sum(1 for s in prs.slides for sh in s.shapes
                  if sh.has_text_frame and sh.text_frame.text.strip())
pics = sum(1 for s in prs.slides for sh in s.shapes if sh.shape_type == 13)
print("editable text shapes:", text_shapes, "| pictures:", pics,
      "| total shapes:", sum(len(s.shapes) for s in prs.slides))
note = prs.slides[0].notes_slide.notes_text_frame.text if prs.slides[0].has_notes_slide else "(无备注页)"
print("page1 notes:", repr(note[:40]))
